"""Build the public observation snapshot and execute exploratory diagnostics.
Usage: python scripts/data/acquire.py --large; python scripts/data/build.py
Inputs are versioned/checksummed in research/data/manifest.json. Raw inputs stay
local. The committed outputs are small derivatives with source-specific licenses.
"""
from pathlib import Path
from datetime import datetime,timezone
import io,json,zipfile,tarfile,hashlib,csv,re,sys
sys.path.insert(0,str(Path(__file__).resolve().parents[2]/"src"))
from collections import Counter
import numpy as np
import openpyxl
import pyshtools as sh
from marswind.observations import temperature_at_depth,threshold_depth,depth_quality,hemisphere_contrast
ROOT=Path(__file__).resolve().parents[2];RAW=ROOT/'data/observations/raw';OUT=ROOT/'research/data'

def clean(x):
    if isinstance(x,np.ndarray):return clean(x.tolist())
    if isinstance(x,dict):return {str(k):clean(v) for k,v in x.items()}
    if isinstance(x,(list,tuple)):return [clean(v) for v in x]
    if isinstance(x,(float,np.floating)):return round(float(x),6) if np.isfinite(x) else None
    if isinstance(x,np.integer):return int(x)
    if isinstance(x,np.bool_):return bool(x)
    return x

def save(name,data):
    p=OUT/name;p.write_text(json.dumps(clean(data),ensure_ascii=False,allow_nan=False,separators=(',',':'))+'\n');print('Saved',name,p.stat().st_size,flush=True)

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def extract_crust():
    folder=RAW/'crustal_models/selected';folder.mkdir(exist_ok=True)
    needed=['Readme.txt','constant-density-summary.xlsx','dichotomy-summary.xlsx','dichotomy_coordinates-JAH-0-360.txt']+[f'Mars-thick-Khan2022-39-{rho}-2900.dat' for rho in [2600,2700,2800,2900]]
    if all((folder/n).exists() for n in needed):return folder
    with tarfile.open(RAW/'crustal_models/InSight-Crustal-Thickness-Archive.tar.gz') as archive:
        for item in archive:
            if item.isfile() and Path(item.name).name in needed:
                (folder/Path(item.name).name).write_bytes(archive.extractfile(item).read())
    return folder

def meteorites():
    with zipfile.ZipFile(RAW/'herd2024/supplementary.zip') as outer:
        z=next(n for n in outer.namelist() if n.endswith('data_s1_to_s3.zip'))
        with zipfile.ZipFile(io.BytesIO(outer.read(z))) as inner:
            name=next(n for n in inner.namelist() if n.endswith('.xlsx') and not n.startswith('__MACOSX'))
            book=openpyxl.load_workbook(io.BytesIO(inner.read(name)),read_only=True,data_only=True)
    samples=[];craters=[];groups=[]
    keys=['name','type','ejection_group','exposure_Ma','exposure_sigma_Ma','terrestrial_kyr','terrestrial_sigma_kyr','ejection_Ma','ejection_sigma_Ma','crystallization_Ma','crystallization_2sigma_Ma']
    for i,row in enumerate(book['Data S1'].iter_rows(min_row=5,max_row=98,max_col=11,values_only=True),5):
        if row[0] and row[1]:samples.append(dict(zip(keys,row),source='Herd et al. 2024, Data S1',row=i))
    keys=['name','id','lat','lon','diameter_km','reported_unit','preservation','secondaries','prior_candidate','model_crater_age_Ma','recurrence_or_literature_age_Ma','terrain_age_Ma']
    for i,row in enumerate(book['Data S2'].iter_rows(min_row=4,max_row=72,max_col=14,values_only=True),4):
        if row[1]:
            c=dict(zip(keys,row[:12]));c['name']=re.sub(r'^\d+\.\s*','',c['name']);c.update(excavation=[],row=i,source='Herd et al. 2024, Data S2',source_url='https://doi.org/10.1126/sciadv.adn2378',status='Candidate, not confirmed');craters.append(c)
        if row[12] and craters:craters[-1]['excavation'].append({'sample':row[12],'maximum_depth_as_reported':row[13]})
    keys=['group','exposure_Ma_as_reported','n_as_reported','crystallization_Ma_as_reported','petrology','main_samples','candidate','preferred_candidate','province']
    for i,row in enumerate(book['Data S3'].iter_rows(min_row=4,max_row=28,max_col=9,values_only=True),4):
        if row[0]:
            g=dict(zip(keys,row));g['alternatives']=[];g['row']=i;g.pop('candidate');groups.append(g)
        if row[6]:groups[-1]['alternatives'].append(row[6])
    craters.append(dict(name='Karratha',id='lagain-karratha',lat=-15.7,lon=203.6,diameter_km=10,source='Lagain et al. 2022, Supplementary Table 2',source_url='https://doi.org/10.1038/s41467-022-31444-8',status='Proposed NWA 7034 source, not confirmed',reported_unit=None,excavation=[],coordinate_note='Table 2 coordinate headings are swapped. Latitude is -15.7; longitude is -156.4 = 203.6 E. Coordinates checked against Table 1 and independent crater coordinates.'))
    for c in craters:
        c['groups']=[g['group'] for g in groups if g['preferred_candidate']==c['name'] or (c['name']=='Karratha' and g['group']=='group 5-2')]
    original=openpyxl.load_workbook(RAW/'ejection_ages/Martian_meteorite_ejection_ages_v1_0.xlsx',data_only=True,read_only=True).active
    age_rows=[list(r) for r in original.iter_rows(min_row=2,max_col=12,values_only=True) if r[0] and r[1]]
    isotopes=openpyxl.load_workbook(RAW/'lagain2022/41467_2022_31444_MOESM3_ESM.xlsx',data_only=True,read_only=True).active
    isotope_rows=[{'row':i,'values':list(r)} for i,r in enumerate(isotopes.iter_rows(max_col=13,values_only=True),1) if any(v is not None for v in r)]
    return {'source':'https://doi.org/10.1126/sciadv.adn2378','license':'Herd-derived tables: CC BY-NC 4.0. Other datasets retain their individual terms.','samples':samples,'groups':groups,'craters':craters,'age_compilation':{'source':'https://zenodo.org/records/10028641','columns':['Name','Type','Subtype','TCRE Ma','sigma','Tterr kyr','sigma','Tej Ma','sigma','Data reference','Compilation reference','Nuclides'],'rows':age_rows},'isotope_table':{'source':'https://doi.org/10.1038/s41467-022-31444-8','table':'Supplementary Data 1, MOESM3','rows':isotope_rows},'notes':['Candidate links are interpretations, not measured coordinates of a meteorite source.','S1 crystallization errors are 2 sigma; S3 summary errors are 1 sigma. Values can differ between tables.','Breccia clast ages, lithification, exposure and ejection are distinct events.','Paired stones, aliquots and ejection groups are not independent regions.']}

def geology(lon,lat,craters):
    import shapely
    from shapely.geometry import Polygon,GeometryCollection,Point
    raw=json.loads((RAW/'geology/units.json').read_text())
    assert raw.get('spatialReference',{}).get('wkid')==104971, 'Expected Mars 2000 Sphere, not terrestrial coordinates'
    assert len(raw['features'])==1311 and not raw.get('exceededTransferLimit'),'Incomplete geologic map'
    units={};features=[];geoms=[]
    for f in raw['features']:
        a=f['attributes'];units[a['Unit']]=a
        # ESRI rings use an even/odd interior rule. XOR retains holes AND islands.
        geom=GeometryCollection()
        for ring in f['geometry']['rings']:
            poly=shapely.make_valid(Polygon(ring));geom=shapely.symmetric_difference(geom,poly)
        geoms.append(geom);features.append(a)
    tree=shapely.STRtree(geoms);names=sorted(units);indices={n:i for i,n in enumerate(names)}
    x,y=np.meshgrid((lon+180)%360-180,lat);points=shapely.points(x.ravel(),y.ravel())
    pairs=tree.query(points,predicate='within');grid=np.full(x.size,-1,dtype=int);hits=np.bincount(pairs[0],minlength=x.size)
    for point_idx,feature_idx in pairs.T:grid[point_idx]=indices[features[feature_idx]['Unit']]
    grid[hits!=1]=-1
    for c in craters:
        matches=tree.query(Point((c['lon']+180)%360-180,c['lat']),predicate='within')
        c['geologic_units_at_point']=[{k:features[m].get(k) for k in ['Unit','UnitDesc','UnitGroup','Interpretation']} for m in matches]
    return {'source':'https://doi.org/10.3133/sim3292','features':len(features),'grid_unassigned_or_ambiguous':int(np.sum(hits!=1)),'units':[{k:units[n].get(k) for k in ['Unit','UnitDesc','UnitGroup','PrimaryCharacteristics','AdditionalCharacteristics','Interpretation']} for n in names],'codes':names,'grid':grid.reshape(x.shape)}

def laboratory():
    path=RAW/'rock_magnetism/magic_contribution_19658.txt'
    tables={}
    for chunk in path.read_text().split('>>>>>>>>>>'):
        lines=chunk.strip().splitlines()
        if len(lines)>2:
            name=lines[0].split('\t')[-1]
            tables[name]=list(csv.DictReader(io.StringIO('\n'.join(lines[1:])),delimiter='\t'))
    return {'source':'https://earthref.org/MagIC/19658','paper':'https://doi.org/10.1029/2022JE007464','license':'CC BY 4.0','interpretation':'Published authors identify terrestrial hand-magnet contamination in the nine paired stones. No ancient paleofield is inferred here.','specimens':tables['specimens'],'measurements':tables['measurements'],'sites':tables['sites'],'notes':['Coordinates in the source are terrestrial find locations. Do not join them to the Mars map.','NRM and laboratory-imparted remanence must not be conflated; method codes and experiment IDs are retained.']}

def main():
    manifest=json.loads((OUT/'manifest.json').read_text());met=meteorites();print('Meteorite tables parsed',flush=True)
    # Two-degree cell centers. MOLA cells are 0.25 degrees; display is a block mean.
    lat=np.arange(-89,90,2);lon=np.arange(1,360,2)
    mola=np.fromfile(RAW/'mola/megt90n000cb.img',dtype='>i2').reshape(720,1440)
    topo=mola.reshape(90,8,180,8).mean(axis=(1,3))[::-1]/1000
    model=sh.SHMagCoeffs.from_file(str(RAW/'magnetic_field/Langlais2019.sh.gz'),lmax=134,skip=4,r0=3393.5e3,header=False,file_units='nT',units='nT',encoding='utf-8')
    x,y=np.meshgrid(lon,lat);fields={}
    for altitude in [150,400]:
        v=model.expand(lat=y.ravel(),lon=x.ravel(),r=np.full(x.size,(3393.5+altitude)*1000))
        fields[str(altitude)]=np.linalg.norm(v,axis=1).reshape(y.shape)
        for c in met['craters']:
            c[f'field_{altitude}_nT']=float(np.linalg.norm(model.expand(lat=float(c['lat']),lon=float(c['lon']),r=(3393.5+altitude)*1000)))
        print('Magnetic field evaluated at',altitude,flush=True)
    folder=extract_crust();crust={}
    for rho in [2600,2700,2800,2900]:
        grid=np.loadtxt(folder/f'Mars-thick-Khan2022-39-{rho}-2900.dat');assert grid.shape==(721,1441)
        # Source grid is node registered (unlike MOLA). Select matching 1+2n deg nodes.
        crust[str(rho)]=grid[np.ix_(((90-lat)*4).astype(int),(lon*4).astype(int))]
        for c in met['craters']:c[f'crust_{rho}_km']=float(grid[round((90-c['lat'])*4),round(c['lon']*4)])
    boundary=np.loadtxt(folder/'dichotomy_coordinates-JAH-0-360.txt')
    # Use the boundary and spherical polygon method supplied by the crust authors.
    dhmask=sh.backends.shtools.Curve2Mask(180,boundary[:,[1,0]],0,sampling=2,extend=True)
    south=dhmask[np.ix_((90-lat).astype(int),lon.astype(int))].astype(bool)
    from marswind.observations import weighted_mean
    crust_contrast=[]
    for density,grid in crust.items():
        n=weighted_mean(grid,lat,~south);ss=weighted_mean(grid,lat,south)
        crust_contrast.append({'south_density_kg_m3':int(density),'north_density_kg_m3':2900,'north_mean_km':n,'south_mean_km':ss,'south_minus_north_km':ss-n})
    boundary_contrast=[]
    for h,grid in fields.items():
        n=weighted_mean(grid,lat,~south);ss=weighted_mean(grid,lat,south)
        boundary_contrast.append({'altitude_km':int(h),'north_mean_nT':n,'south_mean_nT':ss,'south_north_ratio':ss/n})
    print('Crust grids parsed',flush=True)
    geo=geology(lon,lat,met['craters']);print('Geologic map joined',flush=True)
    with zipfile.ZipFile(RAW/'magnetic_depth/MarsMagnetizationDepth.zip') as z:
        def read(suffix):return np.loadtxt(io.BytesIO(z.read('MarsMagnetizationDepth/20_17_8_134_150'+suffix+'.dat')))
        best,lower,upper=read(''),read('_lower'),read('_upper')
    assert np.array_equal(best[:,:2],lower[:,:2]) and np.array_equal(best[:,:2],upper[:,:2])
    bounded,usable=depth_quality(best[:,3],lower[:,3],upper[:,3]);depths=[]
    for i,r in enumerate(best):
        depths.append({'lat':r[0],'lon':r[1]%360,'depth_km':r[3],'lower_km':lower[i,3] if bounded[i] else None,'upper_km':upper[i,3] if bounded[i] else None,'interval_available':bounded[i],'usable':usable[i],'negative_best_fit':r[3]<0,'reduced_chi2':r[5],'source_cap_radius_km':r[2],'window_radius_deg':20,'region':'South' if dhmask[round(90-r[0]),round(r[1]%360)] else 'North'})
    thermal={};ordering={'Pyrrhotite':598.15,'Magnetite':853.15,'Hematite':943.15}
    for hemi in ['North','South']:
        thermal[hemi]={}
        for variant in ['BestModel','Min','Max']:
            radius=np.loadtxt(RAW/f'thermal_profiles/{hemi}_radius_{variant}.csv');temperature=np.loadtxt(RAW/f'thermal_profiles/{hemi}_TemperatureProfile_{variant}.csv')
            d=radius.max()-radius;i=np.argsort(d)
            thermal[hemi][variant]={'depth_km':d[i],'temperature_k':temperature[i],'ordering_depth_km':{name:threshold_depth(radius,temperature,T) for name,T in ordering.items()}}
        for point in depths:
            if point['region']==hemi and point['usable']:
                point['temperature_at_equivalent_depth_k']=float(temperature_at_depth(np.loadtxt(RAW/f'thermal_profiles/{hemi}_radius_BestModel.csv'),np.loadtxt(RAW/f'thermal_profiles/{hemi}_TemperatureProfile_BestModel.csv'),point['depth_km']))
    # Descriptive latitude splits, explicitly not the irregular geologic dichotomy.
    contrasts=[{'altitude_km':h,'exclude_equatorial_degrees':cut,**hemisphere_contrast(fields[str(h)],lat,cut)} for h in [150,400] for cut in [0,20]]
    audit={'windows':len(best),'unavailable_intervals':int((~bounded).sum()),'negative_best_fits':int((best[:,3]<0).sum()),'usable':int(usable.sum()),'negative_and_unavailable':int(((~bounded)&(best[:,3]<0)).sum())}
    retention=[]
    for hemi in ['North','South']:
        pts=[p for p in depths if p['usable'] and p['region']==hemi]
        for carrier,T in ordering.items():
            cool=[p for p in pts if p['temperature_at_equivalent_depth_k']<T]
            retention.append({'hemisphere':hemi,'carrier':carrier,'ordering_temperature_K':T,'usable_windows':len(pts),'below_ordering_temperature':len(cool),'whole_physical_interval_below':sum(p['upper_km']<thermal[hemi]['BestModel']['ordering_depth_km'][carrier] for p in pts),'whole_physical_interval_above':sum(max(0,p['lower_km'])>=thermal[hemi]['BestModel']['ordering_depth_km'][carrier] for p in pts),'best_model_crossing_km':thermal[hemi]['BestModel']['ordering_depth_km'][carrier]})
    results={'generated_at':datetime.now(timezone.utc).isoformat(),'status':'Exploratory diagnostics completed; joint origin inference not performed','magnetic_contrast':contrasts,'boundary_magnetic_contrast':boundary_contrast,'crust_density_sensitivity':crust_contrast,'depth_quality':audit,'thermal_gate':retention,'thermal_note':'A necessary present-day compatibility check using best-fit equivalent depths and hemisphere model endpoints. Not a history of remanence acquisition/survival; source depth is not the bottom of a magnetic layer. No posterior probability or significance test.','south_min_max_identical':sha(RAW/'thermal_profiles/South_TemperatureProfile_Min.csv')==sha(RAW/'thermal_profiles/South_TemperatureProfile_Max.csv'),'counts':{'meteorites':len(met['samples']),'ejection_groups':len(met['groups']),'candidate_craters':len(met['craters']),'geologic_polygons':geo['features'],'geologic_units':len(geo['units']),'downloaded_products':sum(f['status']=='downloaded' for s in manifest for f in s.get('files',[]))},'input_manifest_sha256':sha(OUT/'manifest.json'),'pipeline_sha256':sha(Path(__file__)),'observations_module_sha256':sha(ROOT/'src/marswind/observations.py'),'software':{'numpy':np.__version__,'pyshtools':sh.__version__,'openpyxl':openpyxl.__version__}}
    lab=laboratory();results['counts']['laboratory_measurements']=len(lab['measurements']);results['counts']['laboratory_specimens']=len(lab['specimens']);save('laboratory.json',lab)
    save('results.json',results);save('meteorites.json',met);save('depths.json',depths);save('thermal.json',thermal)
    save('atlas.json',{'latitude':lat,'longitude':lon,'topography_km':np.round(topo,3),'magnetic_nT':{k:np.round(v,3) for k,v in fields.items()},'crust_km':{k:np.round(v,3) for k,v in crust.items()},'geology':geo,'boundary':boundary[::4],'grid_note':'Two-degree cell centers; MOLA 4ppd block means, SH field at centers (full degree 134), crust grids sampled at centers, geologic polygons classified at centers. Display grid is not source resolution. Longitudes east-positive.'})
    figure(lat,lon,topo,fields,depths,thermal,met)
    print(json.dumps(clean(results),indent=2),flush=True)

def figure(lat,lon,topo,fields,depths,thermal,met):
    import matplotlib;matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.colors import LogNorm,LinearSegmentedColormap
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'svg.hashsalt':'mars-first-diagnostics'})
    fig,ax=plt.subplots(2,2,figsize=(13,8),layout='constrained')
    mars_colors=LinearSegmentedColormap.from_list('Mars',['#41465b','#8d929a','#c8b8a5','#c19170','#9a5133','#ce9264','#efe0c5'])
    im=ax[0,0].pcolormesh(lon,lat,topo,cmap=mars_colors,vmin=-8,vmax=16,rasterized=True);fig.colorbar(im,ax=ax[0,0],label='Elevation above areoid (km)')
    ax[0,0].scatter([c['lon'] for c in met['craters']],[c['lat'] for c in met['craters']],c='black',s=14,marker='x');ax[0,0].set_title('MOLA topography + candidate craters (not confirmed sources)')
    im=ax[0,1].pcolormesh(lon,lat,fields['150'],cmap='magma',norm=LogNorm(vmin=1,vmax=1000),rasterized=True);fig.colorbar(im,ax=ax[0,1],extend='both',label='Crustal model |B| at 150 km (nT)');ax[0,1].set_title('MGS/MAVEN: Langlais et al. (2019), degree 134')
    for a in ax[0]:a.set(xlabel='East longitude (°)',ylabel='Planetocentric latitude (°)',xlim=(0,360),ylim=(-90,90))
    usable=[p for p in depths if p['usable']];ax[1,0].errorbar([p['lat'] for p in usable],[p['depth_km'] for p in usable],yerr=[[p['depth_km']-p['lower_km'] for p in usable],[p['upper_km']-p['depth_km'] for p in usable]],fmt='.',ms=3,alpha=.25,color='#9b482c',elinewidth=.5);ax[1,0].set(xlabel='Latitude (°)',ylabel='Equivalent source depth (km)',title='Gong et al. (2021): nonnegative, bounded fits only');ax[1,0].invert_yaxis()
    for h,color in [('North','#486282'),('South','#b64f30')]:
        p=thermal[h]['BestModel'];ax[1,1].plot(p['temperature_k'],p['depth_km'],label=h,color=color)
    for name,T in [('Pyrrhotite',598.15),('Magnetite',853.15),('Hematite',943.15)]:ax[1,1].axvline(T,ls=':',color='gray',lw=.7);ax[1,1].text(T,160,name,rotation=90,va='bottom',fontsize=8)
    ax[1,1].set(xlim=(200,1300),ylim=(200,0),xlabel='Model temperature (K)',ylabel='Depth below local surface (km)',title='Thiriet et al. (2018): present-day endpoints only');ax[1,1].legend()
    fig.suptitle('Mars: first observational consistency diagnostics — no origin mechanism selected',fontsize=14)
    fig.savefig(OUT/'first_diagnostics.png',dpi=180);fig.savefig(OUT/'first_diagnostics.svg',metadata={'Date':None});plt.close(fig)
    svg=OUT/'first_diagnostics.svg';svg.write_text('\n'.join(line.rstrip() for line in svg.read_text().splitlines())+'\n')

if __name__=='__main__':main()
